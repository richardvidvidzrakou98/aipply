"""Tests for cover letter generator module."""

import os

import pytest
from pathlib import Path
from unittest.mock import patch, MagicMock

from docx import Document


@pytest.fixture
def mock_openai():
    """Mock the shared LLM client before CoverLetterGenerator is instantiated."""
    with patch.dict(os.environ, {"OPENROUTER_API_KEY": "test-key-for-unit-tests"}):
        with patch("src.cover_letter_gen.LLMClient") as mock_cls:
            mock_client = MagicMock()
            mock_client.model = "test-model"
            mock_client.temperature = 0.7
            mock_client.available.return_value = True
            mock_cls.return_value = mock_client
            yield mock_client


@pytest.fixture
def cover_gen(mock_openai):
    """Create a CoverLetterGenerator with test config and mocked OpenAI."""
    from src.cover_letter_gen import CoverLetterGenerator

    config = {
        "openai": {
            "model": "gpt-4",
            "temperature": 0.7,
        }
    }
    return CoverLetterGenerator(config=config)


@pytest.fixture
def cover_gen_no_config(mock_openai):
    """Create a CoverLetterGenerator with no config."""
    from src.cover_letter_gen import CoverLetterGenerator

    return CoverLetterGenerator()


@pytest.fixture
def candidate_profile():
    """Sample candidate profile for testing."""
    return {
        "name": "Jane Doe",
        "email": "jane@example.com",
        "location": "San Francisco",
        "summary": "Experienced compliance professional.",
        "strengths": ["Regulatory compliance", "Risk management"],
        "target_roles": ["Compliance Manager", "Risk Analyst"],
    }


class TestCoverLetterGenInit:
    def test_init_with_config(self, cover_gen):
        assert cover_gen.config is not None
        assert "openai" in cover_gen.config
        assert cover_gen.model == "test-model"

    def test_init_default_config(self, cover_gen_no_config):
        assert cover_gen_no_config.config == {}
        assert cover_gen_no_config.model  # shared LLM default

    def test_init_none_config(self, mock_openai):
        from src.cover_letter_gen import CoverLetterGenerator

        g = CoverLetterGenerator(config=None)
        assert g.config == {}


class TestGenerate:
    def test_generate_returns_string(self, cover_gen, mock_openai, candidate_profile):
        """generate should return cover letter text."""
        mock_openai.chat.return_value = (
            "Dear Hiring Manager,\n\n"
            "I am excited about the Compliance Manager position at Acme Corp. "
            "With my background in regulatory compliance and risk management, "
            "I believe I would be a strong fit.\n\n"
            "Sincerely,\nJane Doe"
        )

        result = cover_gen.generate(
            "Seeking a compliance manager with 5 years experience",
            candidate_profile,
            company="Acme Corp",
            role="Compliance Manager",
        )
        assert isinstance(result, str)
        assert "Compliance Manager" in result or "compliance" in result.lower()
        mock_openai.chat.assert_called_once()

    def test_generate_with_empty_description(self, cover_gen, mock_openai, candidate_profile):
        """generate should handle empty job description gracefully."""
        mock_openai.chat.return_value = "Dear Hiring Manager,\n\nI am interested.\n\nSincerely,\nJane"

        result = cover_gen.generate("", candidate_profile)
        assert isinstance(result, str)
        assert len(result) > 0

    def test_generate_with_empty_profile(self, cover_gen, mock_openai):
        """generate should handle empty candidate profile."""
        mock_openai.chat.return_value = "Dear Hiring Manager,\n\nCover letter.\n\nBest regards"

        result = cover_gen.generate("Job description text", {})
        assert isinstance(result, str)

    def test_generate_with_example_letter(self, cover_gen, mock_openai, candidate_profile, tmp_path):
        """generate should incorporate example letter when provided."""
        doc = Document()
        doc.add_paragraph("Example cover letter with professional tone.")
        example_path = tmp_path / "example.docx"
        doc.save(str(example_path))

        mock_openai.chat.return_value = "Generated letter matching style."

        result = cover_gen.generate(
            "Job desc", candidate_profile,
            company="Corp", role="Manager",
            example_letter_path=example_path,
        )
        assert isinstance(result, str)
        mock_openai.chat.assert_called_once()


class TestSaveCoverLetter:
    def test_save_cover_letter_creates_docx(self, cover_gen, tmp_path):
        """save_cover_letter should create a valid .docx file."""
        output_path = tmp_path / "cover_letter.docx"
        content = "Dear Hiring Manager,\n\nI am writing to apply.\n\nSincerely,\nJane Doe"

        result = cover_gen.save_cover_letter(content, output_path)
        assert result == output_path
        assert output_path.exists()

        doc = Document(str(output_path))
        assert len(doc.paragraphs) >= 1

    def test_save_creates_parent_dirs(self, cover_gen, tmp_path):
        """save_cover_letter should create parent directories."""
        output_path = tmp_path / "subdir" / "deep" / "cover_letter.docx"
        content = "Cover letter content"

        result = cover_gen.save_cover_letter(content, output_path)
        assert output_path.exists()


class TestGenerateAndSave:
    def test_generate_and_save_end_to_end(self, cover_gen, mock_openai, candidate_profile, tmp_path):
        """End-to-end: generate_and_save produces a .docx file."""
        output_dir = tmp_path / "output"
        output_dir.mkdir()

        mock_openai.chat.return_value = (
            "Dear Hiring Manager,\n\n"
            "I am excited about this opportunity.\n\n"
            "Sincerely,\nJane Doe"
        )

        result = cover_gen.generate_and_save(
            job_description="Compliance manager role",
            candidate_profile=candidate_profile,
            company="Acme Corp",
            role="Compliance Manager",
            output_dir=output_dir,
        )
        assert Path(result).exists()
        assert str(result).endswith(".docx")

        doc = Document(str(result))
        assert len(doc.paragraphs) >= 1

    def test_generate_and_save_default_output_dir(self, cover_gen, mock_openai, candidate_profile):
        """generate_and_save should create default output dir when none specified."""
        mock_openai.chat.return_value = "Generated letter."

        result = cover_gen.generate_and_save(
            job_description="Job desc",
            candidate_profile=candidate_profile,
            company="TestCo",
            role="Analyst",
        )
        assert Path(result).exists()
        # Clean up
        Path(result).unlink(missing_ok=True)
