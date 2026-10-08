"""Tests for the CV optimizer (scoring + optimization)."""
from unittest.mock import MagicMock, patch

import pytest
from docx import Document

from src.cv_optimizer import CVOptimizer, _heuristic_score


@pytest.fixture
def resume_file(tmp_path):
    doc = Document()
    doc.add_paragraph("Richard Vidzrakou")
    doc.add_paragraph("Professional Summary")
    doc.add_paragraph("DevOps engineer with AWS and Kubernetes experience.")
    doc.add_paragraph("Core Competencies")
    doc.add_paragraph("- Cloud & DevOps: AWS, Docker, Kubernetes, Terraform")
    doc.add_paragraph("- CI/CD: GitHub Actions, Jenkins")
    path = tmp_path / "resume.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def job_description():
    return (
        "We are looking for a DevOps Engineer with strong Kubernetes and "
        "Terraform experience. You will manage CI/CD pipelines on AWS, "
        "automate infrastructure, and improve observability with CloudWatch."
    )


@pytest.fixture
def optimizer_no_llm():
    with patch("src.cv_optimizer.LLMClient") as mock_cls:
        mock_llm = MagicMock()
        mock_llm.available.return_value = False
        mock_llm.model = "heuristic"
        mock_cls.return_value = mock_llm
        yield mock_cls


class TestHeuristicScore:
    def test_returns_score(self, job_description, resume_file):
        from src.resume_tailor import ResumeTailor

        text = ResumeTailor.read_base_resume(resume_file)
        result = _heuristic_score(job_description, text)
        assert 0 <= result["score"] <= 100
        assert "notes" in result

    def test_empty_description(self):
        result = _heuristic_score("", "some resume text")
        assert result["score"] == 0


class TestCVOptimizerNoLLM:
    def test_score_degrades_to_heuristic(self, optimizer_no_llm, resume_file, job_description):
        opt = CVOptimizer()
        result = opt.score_resume(resume_file, job_description)
        assert result["model"] == "heuristic"
        assert isinstance(result["score"], int)

    def test_optimize_returns_safe_fallback_without_llm(self, optimizer_no_llm, resume_file, job_description):
        opt = CVOptimizer()
        result = opt.optimize_resume(resume_file, job_description)
        assert result["summary"].startswith("DevOps engineer")
        assert result["competencies"]
        assert result["model"] == "heuristic_fallback"
        assert result["fallback"] is True


class TestCVOptimizerWithLLM:
    @pytest.fixture
    def optimizer(self):
        with patch("src.cv_optimizer.LLMClient") as mock_cls:
            mock_llm = MagicMock()
            mock_llm.available.return_value = True
            mock_llm.model = "test-model"
            mock_cls.return_value = mock_llm
            opt = CVOptimizer()
            opt.llm = mock_llm
            yield opt

    def test_score_with_llm(self, optimizer, resume_file, job_description):
        optimizer.llm.chat_json.return_value = {
            "requirements": [
                {"requirement": "Kubernetes", "priority": "critical", "category": "Skills", "status": "matched", "evidence": ["Kubernetes experience."]},
                {"requirement": "Terraform", "priority": "required", "category": "Skills", "status": "matched", "evidence": ["Terraform"]},
                {"requirement": "CloudWatch", "priority": "preferred", "category": "Skills", "status": "missing", "evidence": []},
            ],
        }
        result = optimizer.score_resume(resume_file, job_description)
        assert result["score"] == 89
        assert result["requirements"][0]["status"] == "matched"
        assert result["model"] == "test-model"

    def test_unsupported_ai_evidence_becomes_uncertain(self, optimizer, resume_file, job_description):
        optimizer.llm.chat_json.return_value = {"requirements": [
            {"requirement": "Kubernetes", "priority": "critical", "category": "Skills", "status": "matched", "evidence": ["five years of Kubernetes leadership"]},
        ]}
        result = optimizer.score_resume(resume_file, job_description)
        assert result["score"] == 35
        assert result["requirements"][0]["status"] == "uncertain"

    def test_dynamic_missing_requirements_drive_score(self, optimizer, resume_file):
        optimizer.llm.chat_json.return_value = {
            "requirements": [
                {"requirement": "WebRTC protocol expertise", "priority": "critical", "category": "Domain", "status": "missing", "evidence": []},
                {"requirement": "Real-time media pipelines", "priority": "critical", "category": "Experience", "status": "missing", "evidence": []},
                {"requirement": "AWS", "priority": "preferred", "category": "Skills", "status": "matched", "evidence": ["AWS"]},
            ],
        }
        job_description = (
            "Build real-time media pipelines using WebRTC, ICE, DTLS-SRTP, "
            "RTP/RTCP, congestion control, and SDP negotiation. "
            "Write production Go or Rust in asynchronous multi-threaded systems."
        )

        result = optimizer.score_resume(resume_file, job_description)

        assert result["score"] == 9
        assert "WebRTC protocol expertise" in result["critical_missing"]
        assert "Real-time media pipelines" in result["gaps"]

    def test_any_role_specific_requirement_is_scored_without_a_taxonomy(self, optimizer, resume_file, job_description):
        optimizer.llm.chat_json.return_value = {"requirements": [
            {"requirement": "SOC 2 control ownership", "priority": "critical", "category": "Security", "status": "missing", "evidence": []},
            {"requirement": "Quarterly product roadmap", "priority": "preferred", "category": "Product", "status": "missing", "evidence": []},
        ]}

        result = optimizer.score_resume(resume_file, job_description)

        assert result["score"] == 0
        assert [item["requirement"] for item in result["requirements"]] == [
            "SOC 2 control ownership", "Quarterly product roadmap"
        ]

    def test_score_falls_back_when_json_invalid(self, optimizer, resume_file, job_description):
        optimizer.llm.chat_json.return_value = None
        result = optimizer.score_resume(resume_file, job_description)
        assert result["model"] == "heuristic_fallback"

    def test_optimize_with_llm(self, optimizer, resume_file, job_description):
        optimizer.llm.chat_json.return_value = {
            "summary": "DevOps engineer focused on AWS automation.",
            "competencies": ["Cloud & DevOps: AWS, Kubernetes", "CI/CD: GitHub Actions"],
            "changes": ["Reordered competencies", "Rewrote summary"],
        }
        result = optimizer.optimize_resume(resume_file, job_description)
        assert result["summary"].startswith("DevOps engineer")
        assert len(result["competencies"]) == 2
        assert len(result["changes"]) == 2

    def test_optimize_failure_returns_safe_fallback(self, optimizer, resume_file, job_description):
        optimizer.llm.chat_json.return_value = None
        result = optimizer.optimize_resume(resume_file, job_description)
        assert result["model"] == "heuristic_fallback"
        assert result["fallback"] is True
